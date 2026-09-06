"""Builds a downloadable .docx resume tailored to one job, using the
same Profile data the pipeline already parsed -- just re-ordered
projects and a job-specific summary swapped in."""
import io
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _build(profile, tailored_app, doc: Document):
    title = doc.add_heading(profile.name, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links_str = " | ".join(v for v in profile.links.values() if v)
    contact.add_run(f"{profile.location} | {profile.email} | {links_str}").font.size = Pt(10)

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(tailored_app.resume_summary)

    doc.add_heading("Skills", level=1)
    doc.add_paragraph(", ".join(profile.skills))

    doc.add_heading("Projects", level=1)
    projects_by_title = {p["title"]: p for p in profile.projects}
    for title_text in tailored_app.project_order:
        p = projects_by_title.get(title_text)
        if not p:
            continue
        doc.add_paragraph(p["title"]).runs[0].bold = True
        for bullet in p["bullets"]:
            doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("Education", level=1)
    for line in profile.education:
        doc.add_paragraph(line)

    return doc


def build_resume_docx(profile, tailored_app, out_path: str):
    doc = _build(profile, tailored_app, Document())
    doc.save(out_path)
    return out_path


def build_resume_bytes(profile, tailored_app) -> bytes:
    """Same as build_resume_docx but returns bytes in memory -- used by
    the Streamlit app's download_button, which needs bytes, not a path."""
    doc = _build(profile, tailored_app, Document())
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
